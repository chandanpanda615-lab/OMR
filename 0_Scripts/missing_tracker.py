"""The MISSING PDFs LIST (a.k.a. chase list): 5_NetSuite_Booking\\MISSING_PDFS.xlsx, tab "Chase".

One row per invoice we could not book yet. Rows are only ADDED or UPDATED, never wiped by
another script, and a row leaves the list only when its invoice is in _booked.csv
(checked on every save). So nothing falls through the cracks between days.

    download_hul_pdfs.py            GRN row with a blank S3 link        -> "no S3 link (GRN)"
    build_master_workbook.py        GRN row whose hub is not in HUBS    -> "unknown hub ..."
    mark_booked.py                  listed in Master_Entry, not booked  -> "wrong file" (Remark) / "left blank"
    CDMS_Tool\\download_invoices.py  looks every row up in CDMS          -> "CDMS: ..." + CDMS_Recovered.csv

bad_file = the PDF that was wrong. If CDMS / the GRN still offers that same file it is not
listed again; a different file is listed with a "VERIFY" remark.
"""
import csv, os, sys
from datetime import date, datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

OUT_ROOT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "5_NetSuite_Booking")
XLSX = os.path.join(OUT_ROOT, "MISSING_PDFS.xlsx")
RECOVERED = os.path.join(OUT_ROOT, "CDMS_Recovered.csv")   # links CDMS found -> next Master_Entry
TAB = "Chase"
COLS = ["invoice_no", "hub", "brand", "invoice_date", "status", "remark", "bad_file",
        "first_seen", "last_update", "days_waiting"]
WIDTHS = [14, 14, 14, 12, 50, 34, 52, 11, 11, 8]
BOLD = Font(bold=True)
# Remark build_master_workbook puts on a NEW PDF for an invoice whose old PDF was wrong. Left as-is it is
# NOT a "wrong file" remark (mark_booked keeps the old bad_file, so the new PDF is listed again next day).
VERIFY = "was wrong before - VERIFY this new PDF: correct -> type amounts; still wrong -> write your remark"


def has_link(r):
    return (r.get("s3_file_link") or "").strip().startswith("https://cdms-signed-invoice.s3")


def assert_writable(*paths):
    """Excel locks an open workbook. Stop BEFORE changing anything, not half-way through."""
    for p in paths or (XLSX,):
        try:
            if os.path.exists(p):
                open(p, "r+b").close()
        except PermissionError:
            sys.exit(f"Close {os.path.basename(p)} in Excel first, then re-run.")


def iso(v):
    """datetime / 'dd/mm/yyyy' / 'yyyy-mm-dd...' -> 'yyyy-mm-dd' (anything else unchanged)."""
    if isinstance(v, datetime):
        return v.date().isoformat()
    s = str(v or "").strip()
    try:
        return datetime.strptime(s, "%d/%m/%Y").date().isoformat()
    except ValueError:
        return s[:10] if len(s) >= 10 and s[4] == "-" else s


def load():
    """invoice_no -> row dict with every COLS key."""
    if not os.path.exists(XLSX):
        return {}
    wb = load_workbook(XLSX, read_only=True)
    if TAB not in wb.sheetnames:
        return {}
    it = wb[TAB].iter_rows(values_only=True)
    hdr = [str(h) for h in next(it)]
    out = {}
    for r in it:
        d = {c: "" for c in COLS}
        d.update({c: "" if v is None else str(v) for c, v in zip(hdr, r) if c in d})
        if d["invoice_no"].strip():
            out[d["invoice_no"].strip()] = d
    return out


def upsert(chase, inv, **fields):
    """Add inv if new (first_seen = today) and set every non-empty field given."""
    inv = str(inv).strip()
    today = date.today().isoformat()
    row = chase.setdefault(inv, {**{c: "" for c in COLS}, "invoice_no": inv, "first_seen": today})
    for k, v in fields.items():
        if v not in (None, ""):
            row[k] = iso(v) if k == "invoice_date" else str(v).strip()
    row["last_update"] = today
    return row


def save(chase):
    """Write the Chase tab (dropping booked invoices). Returns (path, rows written)."""
    from mark_booked import load_booked
    booked = load_booked()
    today = date.today()
    rows = [r for inv, r in chase.items() if inv not in booked]
    rows.sort(key=lambda r: (r["hub"], r["invoice_date"], r["invoice_no"]))
    os.makedirs(OUT_ROOT, exist_ok=True)
    if os.path.exists(XLSX):
        wb = load_workbook(XLSX)
    else:
        wb = Workbook()
        wb.remove(wb.active)
    for name in (TAB, "Missing_Links", "Wrong_File_Chase"):   # old 2-tab layout is folded into Chase
        if name in wb.sheetnames:
            wb.remove(wb[name])
    ws = wb.create_sheet(TAB, 0)
    ws.append(COLS)
    for c in range(1, len(COLS) + 1):
        ws.cell(1, c).font = BOLD
    for r in rows:
        try:
            r["days_waiting"] = (today - date.fromisoformat(r["first_seen"])).days
        except ValueError:
            r["days_waiting"] = ""
        ws.append([r[c] for c in COLS])
    for i, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(XLSX)
    return XLSX, len(rows)


def grn_rows(csv_path):
    """Rows of the GRN CSV plus CDMS_Recovered.csv (links CDMS found for chase-list invoices),
    so recovered invoices flow through the daily scripts exactly like fresh GRN rows.
    An invoice in both is kept once: the GRN copy if it has a link to a PDF not flagged wrong
    (bad_file), else the recovered copy."""
    with open(csv_path, newline="", encoding="utf-8-sig") as f:
        grn = list(csv.DictReader(f))
    if not os.path.exists(RECOVERED):
        return grn
    with open(RECOVERED, newline="", encoding="utf-8-sig") as f:
        rec = list(csv.DictReader(f))
    bad = {inv: r["bad_file"] for inv, r in load().items() if r["bad_file"]}
    linked = {(r.get("invoice_no") or "").strip() for r in grn if has_link(r)
              and os.path.basename(r["s3_file_link"].strip()) != bad.get((r.get("invoice_no") or "").strip())}
    rec = [dict(r, _src="CDMS") for r in rec if (r.get("invoice_no") or "").strip() not in linked]
    got = {r["invoice_no"].strip() for r in rec}
    return [r for r in grn if (r.get("invoice_no") or "").strip() not in got] + rec


def _selftest():
    assert iso("20/09/2026") == "2026-09-20" and iso("2026-09-20 00:00:00") == "2026-09-20"
    assert iso(datetime(2026, 9, 20, 5)) == "2026-09-20" and iso("") == ""
    c = {}
    upsert(c, " 111 ", hub="OMR", status="no S3 link (GRN)", invoice_date="18/09/2026")
    upsert(c, "111", status="CDMS: PDF not attached yet", remark="")     # blank remark keeps nothing
    assert c["111"]["status"].startswith("CDMS") and c["111"]["hub"] == "OMR"
    assert c["111"]["invoice_date"] == "2026-09-18" and c["111"]["first_seen"] == date.today().isoformat()
    print("selftest OK")


if __name__ == "__main__":
    _selftest()
