import openpyxl
import csv
import os

excel_path = r'C:\Users\chandan.p\Desktop\OMR\OMR_Purchase_Invoices_Entry.xlsx'
header_csv_path = r'C:\Users\chandan.p\Desktop\OMR\1_OMR_Bills_Header.csv'
expenses_csv_path = r'C:\Users\chandan.p\Desktop\OMR\2_OMR_Bills_Expenses.csv'

wb = openpyxl.load_workbook(excel_path, data_only=True)
ws = wb['Invoice_Entry']

# Headers for File 1: Header
header_cols = [
    'External ID',
    'Reference No.',
    'Vendor',
    'Date',
    'Posting Period',
    'Invoice Date',
    'Invoice Receipt Date',
    'Due Date',
    'Approval Status',
    'Currency',
    'Exchange Rate',
    'Department',
    'Class',
    'Location',
    'BRANDS / PROJECTS',
    'Place of Supply',
    'Attached file',
    'Memo'
]

# Headers for File 2: Expenses Sublist (Exact NetSuite Field Names)
expenses_cols = [
    'External ID',
    'Line',
    'Account',
    'Amount',
    'Memo',
    'Department',
    'Class',
    'Location',
    'BRANDS / PROJECTS',
    'INDIA TAX SECTION CODE',
    'INDIA TAX HSN OR SAC CODE',
    'India Tax Nature'
]

header_rows = []
expenses_rows = []

def parse_float(v):
    try:
        return float(v) if v is not None and str(v).strip() != '' else 0.0
    except:
        return 0.0

for row in range(2, ws.max_row + 1):
    inv_no = str(ws.cell(row=row, column=3).value or '').strip()
    if not inv_no:
        continue
    
    inv_date = str(ws.cell(row=row, column=4).value or '').strip()
    pdf_name = str(ws.cell(row=row, column=5).value or '').strip()
    vendor = str(ws.cell(row=row, column=6).value or '').strip()
    location = str(ws.cell(row=row, column=7).value or '').strip()
    brand = str(ws.cell(row=row, column=8).value or '').strip()
    dept = str(ws.cell(row=row, column=9).value or '').strip()
    cls_ = str(ws.cell(row=row, column=10).value or '').strip()
    pos = str(ws.cell(row=row, column=11).value or '').strip()
    memo = str(ws.cell(row=row, column=20).value or '').strip()
    
    taxable_5 = ws.cell(row=row, column=12).value
    taxable_18 = ws.cell(row=row, column=13).value
    taxable_0 = ws.cell(row=row, column=14).value
    
    v5 = parse_float(taxable_5)
    v18 = parse_float(taxable_18)
    v0 = parse_float(taxable_0)
    
    # 1 Header row per invoice
    header_rows.append([
        inv_no,
        inv_no,
        vendor,
        inv_date,
        'Sep 2026',
        inv_date,
        inv_date,
        inv_date,
        'Approved',
        'INR',
        '1.00',
        dept,
        cls_,
        location,
        brand,
        pos,
        pdf_name,
        memo
    ])
    
    # Expense lines
    lines = []
    if v5 > 0:
        lines.append(('4090000', v5))
    if v18 > 0:
        lines.append(('33059011', v18))
    if v0 > 0:
        lines.append(('00000000', v0))
        
    for line_idx, (hsn, amt) in enumerate(lines, start=1):
        expenses_rows.append([
            inv_no,
            line_idx,
            '50020 Cost of Goods Sold : Purchase',
            f'{amt:.2f}',
            memo,
            dept,
            cls_,
            location,
            brand,
            '194Q TDS on Purchases',
            hsn,
            'Goods'
        ])

# Write File 1: Header
with open(header_csv_path, 'w', newline='', encoding='utf-8') as f:
    writer = csv.writer(f)
    writer.writerow(header_cols)
    writer.writerows(header_rows)

# Write File 2: Expenses
with open(expenses_csv_path, 'w', newline='', encoding='utf-8') as f:
    writer = csv.writer(f)
    writer.writerow(expenses_cols)
    writer.writerows(expenses_rows)

print(f'Generated: {header_csv_path} ({len(header_rows)} invoices)')
print(f'Generated: {expenses_csv_path} ({len(expenses_rows)} expense lines)')
