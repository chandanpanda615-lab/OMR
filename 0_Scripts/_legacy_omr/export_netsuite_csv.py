import openpyxl
import csv
import os

excel_path = r'C:\Users\chandan.p\Desktop\OMR\OMR_Purchase_Invoices_Entry.xlsx'
csv_path = r'C:\Users\chandan.p\Desktop\OMR\OMR_NetSuite_Upload.csv'

wb = openpyxl.load_workbook(excel_path, data_only=True)
ws = wb['Invoice_Entry']

# CSV Headers matching NetSuite import fields
csv_headers = [
    'ExternalID',
    'Reference No.',
    'Vendor',
    'Date',
    'Posting Period',
    'Invoice Date',
    'Invoice Receipt Date',
    'Due Date',
    'Approval Status',
    'Currency',
    'Exchange Rate',
    'Department',
    'Class',
    'Location',
    'Brands / Projects',
    'Place of Supply',
    'Attached file',
    'Memo',
    'Expense Account',
    'Expense Amount',
    'Expense Memo',
    'Expense Department',
    'Expense Class',
    'Expense Location',
    'Expense Brands / Projects',
    'Expense India Tax Section Code',
    'Expense India Tax HSN Code',
    'Expense India Tax Nature'
]

rows_to_export = []

# Iterate over rows starting from row 2
for row in range(2, ws.max_row + 1):
    inv_no = str(ws.cell(row=row, column=3).value or '').strip()
    if not inv_no:
        continue
    
    inv_date = str(ws.cell(row=row, column=4).value or '').strip()
    pdf_name = str(ws.cell(row=row, column=5).value or '').strip()
    vendor = str(ws.cell(row=row, column=6).value or '').strip()
    location = str(ws.cell(row=row, column=7).value or '').strip()
    brand = str(ws.cell(row=row, column=8).value or '').strip()
    dept = str(ws.cell(row=row, column=9).value or '').strip()
    cls_ = str(ws.cell(row=row, column=10).value or '').strip()
    pos = str(ws.cell(row=row, column=11).value or '').strip()
    memo = str(ws.cell(row=row, column=20).value or '').strip()
    
    taxable_5 = ws.cell(row=row, column=12).value
    taxable_18 = ws.cell(row=row, column=13).value
    taxable_0 = ws.cell(row=row, column=14).value
    
    # Helper to clean numbers
    def parse_float(v):
        try:
            return float(v) if v is not None and str(v).strip() != '' else 0.0
        except:
            return 0.0

    v5 = parse_float(taxable_5)
    v18 = parse_float(taxable_18)
    v0 = parse_float(taxable_0)
    
    # Check if lines exist
    lines = []
    if v5 > 0:
        lines.append(('4090000', v5))
    if v18 > 0:
        lines.append(('33059011', v18))
    if v0 > 0:
        lines.append(('00000000', v0))
        
    for hsn, amt in lines:
        rows_to_export.append([
            inv_no,                                          # ExternalID
            inv_no,                                          # Reference No.
            vendor,                                          # Vendor
            inv_date,                                        # Date
            'Sep 2026',                                      # Posting Period
            inv_date,                                        # Invoice Date
            inv_date,                                        # Invoice Receipt Date
            inv_date,                                        # Due Date
            'Approved',                                      # Approval Status
            'INR',                                           # Currency
            '1.00',                                          # Exchange Rate
            dept,                                            # Department
            cls_,                                            # Class
            location,                                        # Location
            brand,                                           # Brands / Projects
            pos,                                             # Place of Supply
            pdf_name,                                        # Attached file
            memo,                                            # Memo
            '50020 Cost of Goods Sold : Purchase',           # Expense Account
            f'{amt:.2f}',                                    # Expense Amount
            memo,                                            # Expense Memo
            dept,                                            # Expense Department
            cls_,                                            # Expense Class
            location,                                        # Expense Location
            brand,                                           # Expense Brands / Projects
            '194Q TDS on Purchases',                         # Expense India Tax Section Code
            hsn,                                             # Expense India Tax HSN Code
            'Goods'                                          # Expense India Tax Nature
        ])

with open(csv_path, 'w', newline='', encoding='utf-8') as f:
    writer = csv.writer(f)
    writer.writerow(csv_headers)
    writer.writerows(rows_to_export)

print(f'Successfully generated: {csv_path}')
print(f'Total bill line rows exported: {len(rows_to_export)}')
