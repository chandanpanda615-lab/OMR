import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Invoice_Entry'

# Header columns
headers = [
    'Sr No', 'Folder', 'Invoice Number', 'Invoice Date', 'Attached PDF Filename',
    'Vendor', 'Location', 'Brand / Project', 'Department', 'Class', 'Place of Supply',
    'Taxable 5% (HSN 4090000)', 'Taxable 18% (HSN 33059011)', 'Taxable 0%',
    'Calc GST 5%', 'Calc GST 18%', 'Total Taxable', 'Total GST', 'Invoice Gross Total', 'Memo'
]

ws.append(headers)

# 18 invoice rows from C:\Users\chandan.p\Desktop\OMR
data = [
    ('10th', '9633122377', '10/09/2026', 'OMR-HULS-10092026-9633122377-20260911153446.pdf'),
    ('10th', '9633122378', '10/09/2026', 'OMR-HULS-10092026-9633122378-20260911153458.pdf'),
    ('10th', '9633122379', '10/09/2026', 'OMR-HULS-10092026-9633122379-20260911153516.pdf'),
    ('5th', '9633118905', '05/09/2026', 'OMR-HULS-05092026-9633118905-20260908125719.pdf'),
    ('5th', '9633118906', '05/09/2026', 'OMR-HULS-05092026-9633118906-20260908125707.pdf'),
    ('5th', '9633118907', '05/09/2026', 'OMR-HULS-05092026-9633118907-20260908125731.pdf'),
    ('5th', '9633118994', '05/09/2026', 'OMR-HULS-05092026-9633118994-20260908125756.pdf'),
    ('5th', '9633118995', '05/09/2026', 'OMR-HULS-05092026-9633118995-20260908125808.pdf'),
    ('5th', '9633118996', '05/09/2026', 'OMR-HULS-05092026-9633118996-20260908125823.pdf'),
    ('6th', '9633119242', '06/09/2026', 'OMR-HULS-06092026-9633119242-20260908150126.pdf'),
    ('6th', '9633119243', '06/09/2026', 'OMR-HULS-06092026-9633119243-20260908145901.pdf'),
    ('6th', '9633119244', '06/09/2026', 'OMR-HULS-06092026-9633119244-20260908150142.pdf'),
    ('6th', '9633119272', '06/09/2026', 'OMR-HULS-06092026-9633119272-20260908145922.pdf'),
    ('6th', '9633119273', '06/09/2026', 'OMR-HULS-06092026-9633119273-20260908145938.pdf'),
    ('6th', '9633119274', '06/09/2026', 'OMR-HULS-06092026-9633119274-20260908150153.pdf'),
    ('8th', '9633120170', '08/09/2026', 'OMR-HULS-08092026-9633120170-20260908150206.pdf'),
    ('8th', '9633120171', '08/09/2026', 'OMR-HULS-08092026-9633120171-20260908150218.pdf'),
    ('8th', '9633120172', '08/09/2026', 'OMR-HULS-08092026-9633120172-20260908150229.pdf'),
]

vendor = 'RPP-DIST-0396 HINDUSTAN UNILEVER LIMITED_HUL_OMR'
loc = 'OMR'
brand = 'HUL_OMR'
dept = 'Sales Cost'
cls_ = 'Distribution'
pos = '33-Tamil Nadu'
memo = "Being Purchase for the month of Sep'26"

for idx, (fld, inv, dt, fname) in enumerate(data, start=1):
    r = idx + 1
    # Sample pre-filled verified from earlier inspection
    val_5 = 42830.89 if inv == '9633118905' else (55912.87 if inv == '9633118906' else '')
    val_18 = 0.00 if inv == '9633118905' else (117834.33 if inv == '9633118906' else '')
    val_0 = ''
    
    # Excel formulas
    f_gst5 = f'=IF(L{r}="","",ROUND(L{r}*0.05, 2))'
    f_gst18 = f'=IF(M{r}="","",ROUND(M{r}*0.18, 2))'
    f_taxable = f'=SUM(L{r}:N{r})'
    f_gst = f'=SUM(O{r}:P{r})'
    f_gross = f'=Q{r}+R{r}'
    
    row_vals = [
        idx, fld, inv, dt, fname,
        vendor, loc, brand, dept, cls_, pos,
        val_5, val_18, val_0,
        f_gst5, f_gst18, f_taxable, f_gst, f_gross, memo
    ]
    ws.append(row_vals)

# Styling
header_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')
header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
yellow_fill = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')
calc_fill = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')
border_thin = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

for col in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

for r in range(2, len(data) + 2):
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=r, column=c)
        cell.border = border_thin
        # Input columns: L, M, N (Taxable 5%, 18%, 0%)
        if c in [12, 13, 14]:
            cell.fill = yellow_fill
            cell.number_format = '#,##0.00'
            cell.alignment = Alignment(horizontal='right')
        # Calc columns: O, P, Q, R, S
        elif c in [15, 16, 17, 18, 19]:
            cell.fill = calc_fill
            cell.number_format = '#,##0.00'
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='right')
        elif c in [1, 2, 4]:
            cell.alignment = Alignment(horizontal='center')

# Auto fit column widths
for col in ws.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col:
        val_str = str(cell.value or '')
        if len(val_str) > max_len:
            max_len = len(val_str)
    ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

out_path = r'C:\Users\chandan.p\Desktop\OMR\OMR_Purchase_Invoices_Entry.xlsx'
wb.save(out_path)
print('Successfully saved to:', out_path)
