"""OPTION B: auto-fill Master_Entry.xlsx tax cells from the PDFs, using Gemini.

Flow:  download_hul_pdfs -> build_master_workbook -> Master_Entry.xlsx -> [fill] -> generate_master_csvs
At [fill] you decide per row, BEFORE running this script:
    type the amounts yourself          -> row left alone (your numbers win), no Gemini call
    type x in "No Gemini (type x)"     -> row left alone, no Gemini call
    leave both blank                   -> this script reads it

Every Gemini answer is SAVED FIRST in 5_NetSuite_Booking\\Gemini\\raw\\<md5>.json (one file per
PDF, never overwritten) and the Master is then filled FROM those saved answers. So:
    - the same PDF is never paid for twice (a rebuild or re-run reuses the saved answer, free);
    - build_master_workbook re-applies saved answers whenever it builds a Master;
    - Gemini\\Gemini_Results.xlsx lists every answer next to what was finally put in the import
      file, so you can see how often Gemini was right (refreshed here and by generate_master_csvs).

A PDF with copyable text (Chrompet) is read LOCALLY first, free: saved the same way with model
"pdf-text", but only if every amount passes 5 math checks (see parse_text). Otherwise -> Gemini.

E-INVOICE QR (read locally, free, saved in Gemini\\qr\\<md5>.json), checked on EVERY row, typed or not:
    QR is this invoice      -> "QR Check" OK, Invoice Date = the QR date (= printed invoice date);
                               an answer passes only if taxable + tax = the QR total (within Rs 1)
    QR is another invoice   -> "WRONG FILE - PDF has ...", Remark written (-> Missing PDFs list as
                               wrong file), no Gemini call; generate_master_csvs stops if you type it
    no QR readable (Soukya) -> "NO QR", same checks as before
Gemini first gets only 2-3 pages per invoice (its QR page + the last 2 pages of its part of the
PDF, where Tax Details is); if that answer fails any check the FULL PDF is sent, as before.

Gemini only READS: 5% taxable, 18% taxable, IGST yes/no, printed invoice no (on 2 pages),
invoice date, printed total tax. All tax math stays in Python.
SAFETY: a row is filled ONLY if the computed tax matches the printed total tax (within Rs 1).
Otherwise it stays blank, marked red REVIEW, with the printed tax kept in "Printed Tax (Gemini)" -
and generate_master_csvs refuses a row whose typed amounts don't match that printed tax.

Usage:
    python fill_tax_from_pdfs.py             # needs GEMINI_API_KEY in OMR\\.env (else: saved answers only)
    python fill_tax_from_pdfs.py --view      # only refresh Gemini\\Gemini_Results.xlsx
    python fill_tax_from_pdfs.py --selftest  # offline logic check, no API, no files
"""
import glob
import hashlib
import json
import os
import re
from datetime import datetime
import sys
import time
import urllib.request
from build_entry_workbook import NO_GEMINI, PRINTED_TAX, GEMINI_CHECK, GEMINI_NOTE, QR_CHECK

# ---- pricing / tuning ----
MODEL = "gemini-2.5-flash"
IN_PRICE = 0.30 / 1_000_000
OUT_PRICE = 2.50 / 1_000_000   # output AND thinking tokens
USD_INR = 95.69
THINKING_BUDGET = 1024         # 0 mis-reads the slab split; 512 reads it (verified on real HUL scans); doubled for headroom - it's a cap, not a flat cost
RESCUE_THINKING_BUDGET = 2048   # rescue is a single cropped page for an already-failed PDF - give it more room than the first pass
CHECKSUM_TOLERANCE = 1.0       # rupees; absorbs per-line GST rounding
RETRIES = 3

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
BOOK_ROOT = os.path.join(os.path.dirname(SCRIPT_DIR), "5_NetSuite_Booking")
MASTER = os.path.join(BOOK_ROOT, "Master_Entry.xlsx")
STORE = os.path.join(BOOK_ROOT, "Gemini")
RAW = os.path.join(STORE, "raw")                    # <md5>.json = what Gemini said about that PDF
QR_DIR = os.path.join(STORE, "qr")                  # <md5>.json = the e-invoice QRs in that PDF
VIEW = os.path.join(STORE, "Gemini_Results.xlsx")

COL_DATE, COL_FILE, COL_A5, COL_A18, COL_IGST, COL_HUB = 2, 4, 5, 6, 7, 15

PROMPT = (
    "This is a Hindustan Unilever (HUL) GST tax invoice document, possibly scanned and "
    "rotated/upside-down, and it may bundle SEVERAL separate tax invoices plus logistics "
    "pages (lorry receipt / loading slip). Identify EVERY distinct tax invoice in the file. "
    "For each invoice use ONLY these TWO pages and ignore every other page: "
    "(A) the FIRST page of the invoice, titled 'TAX INVOICE', which prints 'Invoice Number' at the "
    "top and 'Invoice Date' in the right-hand column; "
    "(B) the page that has the 'Tax Details' table (Description / Taxable Value INR / Tax Rate % / "
    "Tax on Net Taxable Turnover) and the totals box (Total Taxes, Invoice Amt); this page also "
    "prints 'Invoice Number' at the top. "
    "In the Tax Details table the GST slabs are printed as: 5% slab = CGST 2.5% + SGST 2.5% (or IGST 5%); "
    "18% slab = CGST 9% + SGST 9% (or IGST 18%). The State GST row repeats the SAME taxable value as its "
    "Central GST row - the slab's taxable value is that value ONCE, never Central + State added. "
    "Amounts use Indian digit grouping (1,75,231.05 = 175231.05). "
    "For each invoice return (read printed values, do not compute): the Invoice Number from page A; "
    "the Invoice Date from page A exactly as printed (dd.mm.yyyy); the Invoice Number from page B; "
    "the total TAXABLE value (pre-tax base) of the 5% slab from page B (0 if none); the total TAXABLE "
    "value of the 18% slab from page B (0 if none); whether IGST is charged (yes/no); and the "
    "'Total Taxes' amount from page B (CGST+SGST, or IGST); the 'Invoice Amt' from the same totals box; "
    "and the page number of page B in this file."
)
RETRY_PROMPT = ("This image is the 'Tax Details' page of a Hindustan Unilever GST tax invoice (it may be rotated). "
                "Read the printed values, do not compute. Amounts use Indian digit grouping (1,17,393.98 = 117393.98).")
RESCUE_KEYS = ("amount_5", "amount_18", "is_igst", "printed_total_tax", "invoice_amt", "dr_cr_adjustments",
               "tax_page_invoice_no")


# ---------- pure logic (unit-tested by --selftest, no API) ----------
def compute_tax(a5, a18):
    """Total GST from the two slab bases. Same total whether intra (2.5+2.5 / 9+9) or IGST (5 / 18)."""
    return round((a5 or 0) * 0.05 + (a18 or 0) * 0.18, 2)


def reconcile(a5, a18, printed_tax):
    """Return (status, note). status in {PASS, REVIEW}; PASS only if computed==printed within tol."""
    if (a5 or 0) == 0 and (a18 or 0) == 0:
        return "REVIEW", "no taxable amounts read"
    if not printed_tax or printed_tax <= 0:
        return "REVIEW", "no printed tax to verify against"
    diff = round(compute_tax(a5, a18) - printed_tax, 2)
    if abs(diff) <= CHECKSUM_TOLERANCE:
        return "PASS", ""
    return "REVIEW", f"checksum off by {diff:,.2f} (computed {compute_tax(a5, a18):,.2f} vs printed {printed_tax:,.2f})"


def amt_gap(a5, a18, printed_tax, invoice_amt, adj=0):
    """taxable + tax - printed Invoice Amt (same as text check 5); None when Invoice Amt was not read
    (answers saved before 25-Sep). Catches a read that is wrong yet adds up, e.g. CGST+SGST rows added.
    Invoice Amt can be hidden under the hub stamp, so a read of Net Payable (= Invoice Amt + Dr/Cr adj)
    is accepted too (25-Sep 9629078879: adj -8,021.00)."""
    if not invoice_amt:
        return None
    gap = round((a5 or 0) + (a18 or 0) + (printed_tax or 0) - invoice_amt, 2)
    net = round(gap + (adj or 0), 2)
    return net if abs(net) < abs(gap) else gap


def qr_gap(a5, a18, q):
    """taxable + tax (our math) - the QR total (TotInvVal); None without a QR. The QR is read by our
    code, not Gemini, so a misread cannot agree with it by accident (archive: 283/283 OK within Rs 1)."""
    if not q:
        return None
    return round((a5 or 0) + (a18 or 0) + compute_tax(a5, a18) - float(q["TotInvVal"]), 2)


def qr_status(ext, qr):
    """-> (QR data of this invoice or None, text for the QR Check column)."""
    if not qr:
        return None, "NO QR"
    q = qr.get(str(ext).strip())
    if q:
        return q, "OK"
    # ponytail: a bundle whose OWN QR is unreadable but a sibling's is would look wrong too (0 of 351 in
    # the archive); if the PDF really is this invoice, type it, clear the Remark, --accept it in generate
    return None, "WRONG FILE - PDF has " + ", ".join(qr)


def slice_pages(wanted, qr, page_count):
    """Pages Gemini gets first: per wanted invoice its QR page (invoice no + date) and the last 2 pages
    of its part of the PDF (up to the next QR page) - where Tax Details was on 134/144 past invoices.
    None when that is not fewer pages than the whole PDF."""
    starts = sorted({v["page"] for v in qr.values()})
    keep = set()
    for q in wanted:
        p = q["page"]
        end = max(p, next((s - 1 for s in starts if s > p), page_count))
        keep |= {p, max(p, end - 1), end}
    return sorted(keep) if len(keep) < page_count else None


def to_ddmmyyyy(s):
    """'19.09.2026' (or 19/09/2026, 19-09-2026) -> '19/09/2026'; None if it is not a real date."""
    try:
        return datetime.strptime(str(s).strip().replace(".", "/").replace("-", "/"), "%d/%m/%Y").strftime("%d/%m/%Y")
    except ValueError:
        return None


def date_ok(d, fn):
    """Printed invoice date (dd/mm/yyyy) must be 0-10 days before the date in the PDF file name
    (catches misreads like 18/05 for 18/09). No date in the file name -> cannot check -> accept."""
    try:
        gap = (datetime.strptime(_iso_from(fn), "%Y-%m-%d") - datetime.strptime(d, "%d/%m/%Y")).days
    except ValueError:
        return True
    return 0 <= gap <= 10


def match_invoice(external_id, extracted):
    """Find the extracted invoice whose printed number == this row's External ID."""
    key = str(external_id).strip()
    for e in extracted:
        if str(e.get("external_id", "")).strip() == key:
            return e
    return None


def judge(ext, fn, invoices, pages=None, q=None):
    """Decide one Master row from the invoices Gemini read in its PDF -> dict with
    status PASS/REVIEW, note, printed (tax), and a5/a18/igst/date when PASS.
    pages = page count of the PDF: a Tax Details page past the end means Gemini made the answer up.
    q = this invoice's e-invoice QR data: taxable + tax must equal its total."""
    m = match_invoice(ext, invoices)
    if m is None:
        found = ", ".join(sorted({str(e.get("external_id", "")).strip() for e in invoices} - {""})) or "none"
        return dict(status="REVIEW", printed=None,
                    note=f"invoice {ext} not printed in PDF (found: {found}) - wrong file?")
    if str(m.get("tax_page_invoice_no", "")).strip() != str(ext).strip():
        return dict(status="REVIEW", printed=None,
                    note=f"invoice no on Tax Details page ({m.get('tax_page_invoice_no')}) != {ext}")
    tp = m.get("retry_page") or m.get("tax_page")   # a passed retry read a real page
    if pages and tp and tp > pages:                  # 25-Sep 9629079589: "page 8" of a 7-page PDF, amounts doubled
        return dict(status="REVIEW", printed=None,
                    note=f"Gemini says Tax Details is page {tp} but the PDF has only {pages} pages - made-up answer, type it")
    d = to_ddmmyyyy(m.get("invoice_date"))
    dnote = "" if d else "invoice date not read"
    if d and not date_ok(d, fn):
        d, dnote = None, f"date {d} looks wrong, sheet date kept"
    status, note = reconcile(m["amount_5"], m["amount_18"], m["printed_total_tax"])
    gap = amt_gap(m["amount_5"], m["amount_18"], m["printed_total_tax"], m.get("invoice_amt"), m.get("dr_cr_adjustments"))
    if status == "PASS" and gap is not None and abs(gap) > CHECKSUM_TOLERANCE:
        status, note = "REVIEW", (f"taxable + tax is off Invoice Amt {m['invoice_amt']:,.2f} by {gap:,.2f}"
                                  " - amounts may be doubled/swapped")
    g = qr_gap(m["amount_5"], m["amount_18"], q)
    if status == "PASS" and g is not None and abs(g) > CHECKSUM_TOLERANCE:
        status, note = "REVIEW", f"taxable + tax is off the QR total {float(q['TotInvVal']):,.2f} by {g:,.2f}"
    printed = round(m.get("printed_total_tax") or 0, 2) or None
    if status == "PASS":
        return dict(status="PASS", printed=printed, a5=round(m["amount_5"], 2), a18=round(m["amount_18"], 2),
                    igst=bool(m["is_igst"]), date=d,
                    note="; ".join(x for x in ("IGST" if m["is_igst"] else "", dnote) if x))
    return dict(status="REVIEW", printed=printed,
                note=f"{note}; read 5%={m['amount_5']:,.2f} 18%={m['amount_18']:,.2f}" + (f"; {dnote}" if dnote else ""))


def _selftest():
    assert compute_tax(100000, 0) == 5000.0
    assert compute_tax(0, 100000) == 18000.0
    assert reconcile(232635.50, 338581.45, 72576.34)[0] == "PASS"          # real BYTI row
    assert reconcile(151255.70, 0, 7562.84)[0] == "PASS"                    # real OMR row (pure 5%)
    assert reconcile(0, 505700.77, 72576.34)[0] == "REVIEW"                 # wrong split -> caught
    assert reconcile(0, 0, 9815.38)[0] == "REVIEW"                          # nothing read
    assert reconcile(100000, 0, 0)[0] == "REVIEW"                           # no printed tax
    ex = [{"external_id": "9629076836"}, {"external_id": "9629076837"}]
    assert match_invoice("9629076837", ex)["external_id"] == "9629076837"   # bundle match
    assert match_invoice(9629076838, ex) is None                            # not in bundle -> None
    assert to_ddmmyyyy("19.09.2026") == "19/09/2026" and to_ddmmyyyy("31.02.2026") is None and to_ddmmyyyy(None) is None
    f = "TLBL-HUL-19092026-9633128802-20260920100901.pdf"
    assert date_ok("18/09/2026", f) and not date_ok("18/05/2026", f) and not date_ok("20/09/2026", f)   # 18/05 = real misread
    inv = [{"external_id": "9633128802", "tax_page_invoice_no": "9633128802", "invoice_date": "18.09.2026",
            "amount_5": 151255.70, "amount_18": 0, "is_igst": False, "printed_total_tax": 7562.84}]
    j = judge("9633128802", f, inv)
    assert j["status"] == "PASS" and j["a5"] == 151255.70 and j["date"] == "18/09/2026" and j["printed"] == 7562.84
    assert judge("9633128803", f, inv)["note"].endswith("wrong file?")                    # not in this PDF
    inv[0]["tax_page"] = 8
    assert judge("9633128802", f, inv, pages=8)["status"] == "PASS"
    assert "made-up" in judge("9633128802", f, inv, pages=7)["note"]                      # page past the end
    inv[0]["retry_page"] = 6
    assert judge("9633128802", f, inv, pages=7)["status"] == "PASS"                       # retry read a real page
    del inv[0]["tax_page"], inv[0]["retry_page"]
    # 25-Sep 9629079589: CGST+SGST rows added -> doubled amount AND doubled tax still pass the checksum;
    # the printed Invoice Amt (4,49,219.12) catches it
    dbl = [{"external_id": "9629079589", "tax_page_invoice_no": "9629079589", "invoice_date": "24.09.2026",
            "amount_5": 855655.32, "amount_18": 0, "is_igst": False, "printed_total_tax": 42782.92, "invoice_amt": 449219.12}]
    g = "TLBL-HUL-24092026-9629079589-20260925063608.pdf"
    assert "Invoice Amt" in judge("9629079589", g, dbl)["note"]
    dbl[0].update(amount_5=427827.66, printed_total_tax=21391.46)
    assert judge("9629079589", g, dbl)["status"] == "PASS"                                 # real values pass
    assert amt_gap(1, 2, 3, None) is None                                                  # old answers: no check
    assert abs(amt_gap(57043.02, 402787.93, 75353.94, 527163.89, -8021.00)) <= 1               # Net Payable read, adj explains it
    inv[0]["amount_18"] = 1000
    j = judge("9633128802", f, inv)
    assert j["status"] == "REVIEW" and j["printed"] == 7562.84 and "a5" not in j          # checksum off -> not filled
    # text read: real 9633132412 layout (values print BEFORE their labels)
    p1 = "TAX INVOICE\nInvoice Number  :\nPage 1 of 2\n9633132412\n" + "x" * 200 + "\n23.09.2026\nInvoice Date    :\n"
    p2 = ("TAX INVOICE\nInvoice Number  :\nPage 2 of 2\n9633132412\n86,662.68\nTotal Taxes            :\n"
          "7,38,409.61\nInvoice Amt            :\nTax Details\n"
          "IN: Central GST-AA\n2,35,783.66\n2.5\n5,894.6\nIN: Central GST-AA\n4,15,963.27\n9\n37,436.74\n"
          "IN: State GST-AA\n2,35,783.66\n2.5\n5,894.6\nIN: State GST-AA\n4,15,963.27\n9\n37,436.74\n" + "x" * 100)
    inv, why = parse_text([p2, p1])                                                        # any page order
    assert why is None and len(inv) == 1, why
    e = inv[0]
    assert (e["amount_5"], e["amount_18"], e["printed_total_tax"], e["invoice_date"], e["tax_page"], e["is_igst"]) \
        == (235783.66, 415963.27, 86662.68, "23.09.2026", 1, False)
    assert judge("9633132412", "CRMP-HULS-23092026-9633132412-20260923175700.pdf", inv)["status"] == "PASS"
    assert "CGST base != SGST" in parse_text([p1, p2.replace("4,15,963.27\n9\n37,436.74\nIN: State", "4,15,963.72\n9\n37,436.74\nIN: State")])[1]
    assert "Invoice Amt" in parse_text([p1, p2.replace("7,38,409.61", "7,38,490.61")])[1]      # one digit off
    assert "no text" in parse_text([p1, ""])[1]                                              # scanned page -> Gemini
    # e-invoice QR: real 9629078879 (Tumkur) - booked 57,043.02 / 4,02,787.93, QR total 5,35,184.89
    q = {"DocNo": "9629078879", "DocDt": "22/09/2026", "TotInvVal": 535184.89, "page": 1}
    assert abs(qr_gap(57043.02, 402787.93, q)) <= 1 and qr_gap(1, 2, None) is None
    assert qr_status("9629078879", {"9629078879": q}) == (q, "OK")
    assert qr_status("9629078879", {}) == (None, "NO QR")
    assert qr_status("9629076836", {"9629076837": q, "9629076838": q})[1] == "WRONG FILE - PDF has 9629076837, 9629076838"
    tm = [{"external_id": "9629078879", "tax_page_invoice_no": "9629078879", "invoice_date": "22.09.2026",
           "amount_5": 57043.02, "amount_18": 402787.93, "is_igst": False, "printed_total_tax": 75353.94}]
    h = "TMKR-HUL-24092026-9629078879-20260924200317.pdf"
    assert judge("9629078879", h, tm, q=q)["status"] == "PASS"
    tm[0].update(amount_5=402787.93, amount_18=57043.02, printed_total_tax=30406.87)          # slabs swapped, adds up
    assert judge("9629078879", h, tm)["status"] == "PASS" and "QR total" in judge("9629078879", h, tm, q=q)["note"]
    two = {"A": {"page": 1}, "B": {"page": 6}}                                                # bundle: B starts on page 6
    assert slice_pages([two["A"]], two, 10) == [1, 4, 5] and slice_pages([two["B"]], two, 10) == [6, 9, 10]
    assert slice_pages([two["A"], two["B"]], two, 10) == [1, 4, 5, 6, 9, 10]
    assert slice_pages([{"page": 1}], {"A": {"page": 1}}, 3) is None                         # 3 pages: send it all
    print("selftest OK")


# ---------- saved answers (Gemini\raw\<md5>.json) ----------
def _md5(path):
    h = hashlib.md5()
    with open(path, "rb") as f:
        while chunk := f.read(8192):
            h.update(chunk)
    return h.hexdigest()


def load_raw(md5):
    p = os.path.join(RAW, md5 + ".json")
    if not os.path.exists(p):
        return None
    with open(p, encoding="utf-8") as f:
        return json.load(f)


def save_raw(rec):
    """Written right after each paid call, before anything else can fail."""
    os.makedirs(RAW, exist_ok=True)
    p = os.path.join(RAW, rec["md5"] + ".json")
    with open(p + ".tmp", "w", encoding="utf-8") as f:
        json.dump(rec, f, indent=1)
    os.replace(p + ".tmp", p)


# ---------- e-invoice QR: read locally, free (zxing-cpp), never sent anywhere ----------
def read_qr(path, md5):
    """-> {DocNo: QR data + "page"} for every GST e-invoice QR in the PDF; cached in Gemini\\qr."""
    p = os.path.join(QR_DIR, md5 + ".json")
    if os.path.exists(p):
        with open(p, encoding="utf-8") as f:
            return json.load(f)
    import base64
    import io
    import fitz
    import zxingcpp
    from PIL import Image
    found = {}

    def scan(im, page):
        for b in zxingcpp.read_barcodes(im, formats=zxingcpp.BarcodeFormat.QRCode, try_rotate=True, try_downscale=True):
            parts = b.text.split(".")                 # NIC-signed JWT: header.payload.signature
            try:
                pl = json.loads(base64.urlsafe_b64decode(parts[1] + "=" * (-len(parts[1]) % 4)))
                q = json.loads(pl["data"]) if isinstance(pl["data"], str) else pl["data"]
                found.setdefault(str(q["DocNo"]).strip(), dict(q, page=page))
            except (ValueError, KeyError, TypeError, IndexError):
                pass                                  # not an e-invoice QR

    with fitz.open(path) as doc:
        for i, pg in enumerate(doc):
            for x in pg.get_images(full=True):
                try:
                    scan(Image.open(io.BytesIO(doc.extract_image(x[0])["image"])).convert("L"), i + 1)
                except Exception:                     # odd image formats: the render below covers them
                    pass
        if not found:                                 # text PDFs draw the QR as vectors
            for i, pg in enumerate(doc):
                pix = pg.get_pixmap(dpi=300, colorspace=fitz.csGRAY)
                scan(Image.frombytes("L", (pix.width, pix.height), pix.samples), i + 1)
    os.makedirs(QR_DIR, exist_ok=True)
    with open(p, "w", encoding="utf-8") as f:
        json.dump(found, f, indent=1)
    return found


# ---------- text PDFs: read locally, free ----------
_N = r"(-?[\d,]+\.?\d*)"
_TAX_ROW = re.compile(r"IN: (Central|State|Integrated) GST[^\n]*\n\s*" + _N + r"\s*\n\s*" + _N + r"\s*\n\s*" + _N)


def _num(s):
    return float(s.replace(",", ""))


def parse_text(pages):
    """Page texts of one PDF -> (invoice dicts shaped like Gemini's, None) or (None, reason).
    Nothing is accepted unless ALL 5 math checks hold (Rs 1 tolerance; HUL rounds per item line):
    1 row base x rate = row tax  2 CGST base = SGST base  3 row taxes add up to Total Taxes
    4 5%/18% bases give Total Taxes  5 bases + Total Taxes = Invoice Amt."""
    tol, by_inv = CHECKSUM_TOLERANCE, {}
    for i, t in enumerate(pages, 1):
        if len(t.strip()) < 200:
            return None, f"page {i} has no text (scan)"
        m = re.search(r"Page \d+ of \d+\s*\n\s*(\d{10})", t)
        if not m:
            return None, f"page {i}: invoice number not found"
        by_inv.setdefault(m.group(1), []).append((i, t))
    out = []
    for inv, pl in by_inv.items():
        date = tax_page = None
        for i, t in pl:
            m = re.search(r"(\d\d\.\d\d\.\d{4})\s*\nInvoice Date", t)
            date = date or (m and m.group(1))
            if "Tax Details" in t and "Total Taxes" in t:
                tax_page = (i, t)
        if not tax_page:
            return None, f"{inv}: Tax Details page not found"
        k, t = tax_page
        rows = _TAX_ROW.findall(t)
        if not rows:
            return None, f"{inv}: no tax rows"
        base, row_tax = {}, 0.0
        for kind, b, r, tx in rows:
            b, r, tx = _num(b), _num(r), _num(tx)
            if abs(b * r / 100 - tx) > tol:
                return None, f"{inv}: {kind} {r}% row: base x rate != tax"                    # check 1
            base[(kind, r)] = base.get((kind, r), 0) + b
            row_tax += tx
        a5 = a18 = 0.0
        for (kind, r), b in base.items():
            if kind == "State":
                if abs(b - base.get(("Central", r), -1)) > 0.01:
                    return None, f"{inv}: CGST base != SGST base at {r}%"                    # check 2
                continue
            if r in (2.5, 5):
                a5 += b
            elif r in (9, 18):
                a18 += b
            else:
                return None, f"{inv}: unknown tax rate {r}%"
        mt, ma = re.search(_N + r"\s*\n\s*Total Taxes", t), re.search(_N + r"\s*\n\s*Invoice Amt", t)
        if not (mt and ma):
            return None, f"{inv}: Total Taxes / Invoice Amt not found"
        total, amt = _num(mt.group(1)), _num(ma.group(1))
        if abs(row_tax - total) > tol:
            return None, f"{inv}: tax rows {row_tax:,.2f} != Total Taxes {total:,.2f}"          # check 3
        if reconcile(a5, a18, total)[0] != "PASS":
            return None, f"{inv}: 5%/18% tax != Total Taxes {total:,.2f}"                      # check 4
        if abs(a5 + a18 + total - amt) > tol:
            return None, f"{inv}: taxable + tax != Invoice Amt {amt:,.2f}"                     # check 5
        out.append({"external_id": inv, "invoice_date": date, "tax_page_invoice_no": inv,
                    "amount_5": round(a5, 2), "amount_18": round(a18, 2),
                    "is_igst": any(kind == "Integrated" for kind, _ in base),
                    "printed_total_tax": total, "invoice_amt": amt, "tax_page": k})
    return out, None


def read_text(path, md5, hub):
    """Text PDF -> saved record (model pdf-text, Rs 0), or None -> caller uses Gemini."""
    import fitz
    with fitz.open(path) as doc:
        invoices, why = parse_text([pg.get_text() for pg in doc])
    if invoices is None:
        print(f"    text read rejected ({why}) -> Gemini")
        return None
    rec = {"md5": md5, "file": os.path.basename(path), "hub": hub,
           "read_on": datetime.now().strftime("%Y-%m-%d %H:%M"), "model": "pdf-text",
           "cost_rs": 0.0, "invoices": invoices}
    save_raw(rec)
    return rec


# ---------- Gemini I/O ----------
def _extract(client, types, Extract, pdf_path, data=None):
    """One Gemini call -> list of invoice dicts + (in, out, thinking) tokens. Retries on transient errors.
    data = bytes of a few-page PDF cut from pdf_path, sent inline instead of uploading the file."""
    last = None
    for attempt in range(1, RETRIES + 1):
        try:
            up = None if data else client.files.upload(file=pdf_path)
            r = client.models.generate_content(
                model=MODEL,
                contents=[up or types.Part.from_bytes(data=data, mime_type="application/pdf"), PROMPT],
                config=types.GenerateContentConfig(
                    response_mime_type="application/json",
                    response_schema=list[Extract],
                    thinking_config=types.ThinkingConfig(thinking_budget=THINKING_BUDGET),
                ),
            )
            try:
                if up:
                    client.files.delete(name=up.name)
            except Exception:
                pass
            m = r.usage_metadata
            inv = [e.model_dump() for e in (r.parsed or [])]
            return inv, (m.prompt_token_count or 0, m.candidates_token_count or 0,
                         getattr(m, "thoughts_token_count", 0) or 0)
        except Exception as e:
            last = e
            print(f"    [retry {attempt}/{RETRIES}] {os.path.basename(pdf_path)}: {e}")
            time.sleep(2 * attempt)
    raise last


def make_caller(key):
    """-> call(path, md5, hub): asks Gemini about one PDF, SAVES the answer, returns it."""
    from google import genai
    from google.genai import types
    from pydantic import BaseModel, Field

    class Extract(BaseModel):
        external_id: str = Field(description="Invoice Number printed at the top of the TAX INVOICE first page (page A)")
        invoice_date: str = Field(description="Invoice Date from page A, exactly as printed, dd.mm.yyyy")
        tax_page_invoice_no: str = Field(description="Invoice Number printed at the top of the Tax Details page (page B)")
        amount_5: float = Field(description="Taxable Value of the ONE Central GST row with Tax Rate 2.5 (IGST: the row "
                                            "with rate 5). The State GST row repeats the same value - do NOT add them. 0 if none")
        amount_18: float = Field(description="Taxable Value of the ONE Central GST row with Tax Rate 9 (IGST: the row "
                                             "with rate 18). The State GST row repeats the same value - do NOT add them. 0 if none")
        is_igst: bool = Field(description="true if IGST charged (interstate), false if CGST+SGST")
        printed_total_tax: float = Field(description="Total tax printed on the invoice (CGST+SGST or IGST); verification only")
        invoice_amt: float = Field(description="'Invoice Amt' printed in the totals box of page B (NOT 'Net Payable'); verification only")
        dr_cr_adjustments: float = Field(description="'Dr/Cr Adjustments' in the same totals box, with its sign; 0 if none")
        tax_page: int = Field(description="1-based page number, in this file, of page B (the Tax Details page)")

    class TaxPage(BaseModel):
        tax_page_invoice_no: str = Field(description="Invoice Number printed at the top of this page")
        amount_5: float = Field(description="Taxable Value of the ONE Central GST row with Tax Rate 2.5 (IGST: the row "
                                            "with rate 5). The State GST row repeats the same value - do NOT add them. 0 if none")
        amount_18: float = Field(description="Taxable Value of the ONE Central GST row with Tax Rate 9 (IGST: the row "
                                             "with rate 18). The State GST row repeats the same value - do NOT add them. 0 if none")
        is_igst: bool = Field(description="true if IGST, false if CGST+SGST")
        printed_total_tax: float = Field(description="'Total Taxes' in the totals box")
        invoice_amt: float = Field(description="'Invoice Amt' in the totals box (NOT 'Net Payable')")
        dr_cr_adjustments: float = Field(description="'Dr/Cr Adjustments' in the totals box, with its sign; 0 if none")

    client = genai.Client(api_key=key)

    def retry(path, e, q=None):
        """REVIEW rescue: send ONLY the Tax Details page, turned upright, as one image (the whole
        13-page scan confused Gemini; the single upright page read right 3/3 on 9633128588).
        -> (answer that passes, or None; cost_rs). Stops at the first answer that passes (q: its QR too)."""
        import io
        import fitz
        from PIL import Image
        doc = fitz.open(path)
        k = e.get("tax_page") or 0
        k = k if 1 <= k <= doc.page_count else doc.page_count   # older answers: guess the last page
        pix = doc[k - 1].get_pixmap(dpi=150)
        im = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        cost = 0.0
        # the HUL Tax Details page is wide: a tall image is a sideways scan (either way round)
        for rot in ((90, 270) if im.height > im.width else (0, 180)):
            buf = io.BytesIO()
            im.rotate(rot, expand=True).save(buf, "PNG")
            r = client.models.generate_content(
                model=MODEL, contents=[types.Part.from_bytes(data=buf.getvalue(), mime_type="image/png"), RETRY_PROMPT],
                config=types.GenerateContentConfig(
                    response_mime_type="application/json", response_schema=TaxPage,
                    thinking_config=types.ThinkingConfig(thinking_budget=RESCUE_THINKING_BUDGET)))
            m = r.usage_metadata
            cost += ((m.prompt_token_count or 0) * IN_PRICE + ((m.candidates_token_count or 0)
                     + (getattr(m, "thoughts_token_count", 0) or 0)) * OUT_PRICE) * USD_INR
            a = r.parsed.model_dump() if r.parsed else None
            if (a and str(a["tax_page_invoice_no"]).strip() == str(e.get("external_id", "")).strip()
                    and reconcile(a["amount_5"], a["amount_18"], a["printed_total_tax"])[0] == "PASS"
                    and abs(amt_gap(a["amount_5"], a["amount_18"], a["printed_total_tax"], a["invoice_amt"], a["dr_cr_adjustments"]) or 0)
                    <= CHECKSUM_TOLERANCE
                    and abs(qr_gap(a["amount_5"], a["amount_18"], q) or 0) <= CHECKSUM_TOLERANCE):
                return dict(a, page=k, rot=rot), round(cost, 2)
        return None, round(cost, 2)

    def call(path, md5, hub, pages=None):
        """pages (1-based) -> only those pages are sent; the answer's page numbers are mapped back."""
        data = None
        if pages:
            import fitz
            with fitz.open(path) as src, fitz.open() as cut:
                for p in pages:
                    cut.insert_pdf(src, from_page=p - 1, to_page=p - 1)
                data = cut.tobytes()
        invoices, (ti, to, tt) = _extract(client, types, Extract, path, data)
        for e in invoices if pages else ():
            tp = e.get("tax_page")
            # ponytail: a page outside the slice becomes None (page check skipped) - slices only go out
            # with a QR, and the QR total check catches a made-up answer
            e["tax_page"] = pages[tp - 1] if isinstance(tp, int) and 1 <= tp <= len(pages) else None
        rec = {"md5": md5, "file": os.path.basename(path), "hub": hub,
               "read_on": datetime.now().strftime("%Y-%m-%d %H:%M"), "model": MODEL,
               "thinking_budget": THINKING_BUDGET, "tokens": {"in": ti, "out": to, "thinking": tt},
               "cost_rs": round((ti * IN_PRICE + (to + tt) * OUT_PRICE) * USD_INR, 2),
               "invoices": invoices}
        if pages:
            rec["pages_sent"] = pages
        save_raw(rec)
        return rec
    call.retry = retry
    return call


# ---------- fill a sheet from saved answers (used here AND by build_master_workbook) ----------
def _pdf_index(hubs, root):
    """Map every PDF basename under each present hub's PDFs folder (below `root`) -> full path."""
    idx = {}
    for loc in hubs:
        base = os.path.join(root, loc, "PDFs")
        if not os.path.isdir(base):
            continue
        for dp, _, files in os.walk(base):
            for fn in files:
                if fn.lower().endswith(".pdf"):
                    idx.setdefault(fn, os.path.join(dp, fn))
    return idx


def _iso_from(basename):
    """SUKA-HUL-17092026-...pdf -> 2026-09-17 (matches download_from_links folder convention)."""
    t = basename[:-4].split("-")
    if len(t) > 2 and len(t[2]) == 8 and t[2].isdigit():
        d = t[2]
        return f"{d[4:]}-{d[2:4]}-{d[0:2]}"
    return "_fetched"


def _download(url, hub, basename, root):
    """Download the PDF from its column-D link into the hub's PDFs folder under `root`
    (so the generate step's zip finds it too). Cached copies are reused, never re-fetched."""
    folder = os.path.join(root, hub or "_misc", "PDFs", _iso_from(basename))
    os.makedirs(folder, exist_ok=True)
    target = os.path.join(folder, basename)
    if not os.path.exists(target):
        urllib.request.urlretrieve(url, target)
    return target


def gemini_cols(ws):
    """Header name -> column number; adds the Gemini / QR columns if an older sheet lacks them."""
    col = {c.value: c.column for c in ws[1] if c.value}
    for name in (NO_GEMINI, PRINTED_TAX, GEMINI_CHECK, GEMINI_NOTE, QR_CHECK):
        if name not in col:
            col[name] = ws.max_column + 1
            ws.cell(1, col[name]).value = name
    return col


def fill_sheet(ws, root, call=None):
    """Check EVERY row's PDF against its e-invoice QR (free: QR Check, invoice date, WRONG FILE Remark),
    then fill every row that is not typed by you and not marked x, from SAVED answers, then a free
    local text read of the PDF (if it has copyable text and passes the 5 math checks).
    call=None (build_master_workbook, or no API key): saved answers + free text reads - never costs money.
    call given: a PDF with no saved/text answer is sent to Gemini once (2-3 pages per invoice first,
    the full PDF if that fails) and the answer saved first. Returns a dict of counts."""
    from openpyxl.styles import PatternFill
    from missing_tracker import VERIFY
    green, red = PatternFill("solid", fgColor="C6EFCE"), PatternFill("solid", fgColor="FFC7CE")
    col = gemini_cols(ws)
    n = dict(rows=0, marked=0, typed=0, saved=0, text_read=0, new_calls=0, cost_rs=0.0, passed=0, review=0,
             not_read=0, retried=0, rescued=0, qr_ok=0, wrong_file=0, no_qr=0, date_fixed=0, sliced=0,
             full_after_slice=0)
    todo = []
    for row in range(2, ws.max_row + 1):
        ext, fn = ws.cell(row, 1).value, ws.cell(row, COL_FILE).value
        if ext is None or not fn:
            continue
        n["rows"] += 1
        read = True                      # False: marked x / typed by you -> QR check only, amounts untouched
        if str(ws.cell(row, col[NO_GEMINI]).value or "").strip():
            n["marked"] += 1
            read = False
        elif ws.cell(row, COL_A5).value not in (None, "") or ws.cell(row, COL_A18).value not in (None, ""):
            n["typed"] += 1
            read = False
        d = ws.cell(row, COL_FILE)
        todo.append((row, str(ext).strip(), str(fn).strip(), d.hyperlink.target if d.hyperlink else None,
                     str(ws.cell(row, COL_HUB).value or "").strip(), read))
    if not todo:
        return n

    def mark(row, printed, status, note):
        ws.cell(row, col[PRINTED_TAX]).value = printed or None
        c = ws.cell(row, col[GEMINI_CHECK])
        c.value, c.fill = status, (green if status == "PASS" else red)
        ws.cell(row, col[GEMINI_NOTE]).value = note or None

    cache = _pdf_index({t[4] for t in todo if t[4]}, root)
    groups = {}
    for row, ext, fn, url, hub, read in todo:
        path = cache.get(fn)
        if not path and url and call and read:
            try:
                path = cache[fn] = _download(url, hub, fn, root)
            except Exception as e:
                print(f"    [download failed] {fn}: {e}")
        if not path:
            if not read:
                continue
            if call:
                mark(row, None, "REVIEW", "PDF unavailable (no local copy and no/failed link in column D)")
                n["review"] += 1
            else:
                n["not_read"] += 1
            continue
        md5 = _md5(path)
        groups.setdefault(md5, {"path": path, "hub": hub, "rows": []})["rows"].append((row, ext, fn, read))

    import fitz
    for i, (md5, g) in enumerate(groups.items(), 1):
        name = os.path.basename(g["path"])
        try:
            qr = read_qr(g["path"], md5)
        except Exception as e:
            print(f"    [QR read failed] {name}: {e}")
            qr = {}
        with fitz.open(g["path"]) as doc:
            pages = doc.page_count
        need = []                        # rows whose amounts are read below, with their QR (or None)
        for row, ext, fn, read in g["rows"]:
            q, status = qr_status(ext, qr)
            if q:
                n["qr_ok"] += 1
                c = ws.cell(row, COL_DATE)
                old = c.value.strftime("%d/%m/%Y") if hasattr(c.value, "strftime") else str(c.value or "").strip()
                if q.get("DocDt") and old != q["DocDt"]:
                    c.value = q["DocDt"]                 # we book on the invoice date = the QR date (132/132 checked)
                    status += f" - date changed from {old}"
                    n["date_fixed"] += 1
            elif status == "NO QR":
                n["no_qr"] += 1
            else:
                n["wrong_file"] += 1
                rem = col.get("Remark")
                if rem and str(ws.cell(row, rem).value or "").strip() in ("", VERIFY):
                    ws.cell(row, rem).value = status + " (QR)"      # mark_booked -> Missing PDFs list as wrong file
                if read:
                    mark(row, None, "REVIEW", status + " - not sent to Gemini")
                    n["review"] += 1
            qc = ws.cell(row, col[QR_CHECK])
            if not (status == "OK" and str(qc.value or "").startswith("OK")):   # keep "OK - date changed" on re-runs
                qc.value = status
            if read and (q or status == "NO QR"):
                need.append((row, ext, fn, q))
        if not need:
            continue

        rec = load_raw(md5)
        if rec is not None:
            n["saved"] += len(need)
        else:
            try:
                rec = read_text(g["path"], md5, g["hub"])       # free; None -> Gemini below
            except Exception as e:
                print(f"    [text read failed] {name}: {e} -> Gemini")
            if rec is not None:
                n["text_read"] += len(need)
                print(f"  [{i}/{len(groups)}] read from PDF text, free: {name}")
        if rec is not None:
            pass
        elif call is None:
            n["not_read"] += len(need)
            continue
        else:
            cut = slice_pages([q for *_, q in need], qr, pages) if all(q for *_, q in need) else None
            print(f"  [{i}/{len(groups)}] Gemini reads {name}  ({len(need)} row(s))"
                  + (f"  - pages {cut} of {pages}" if cut else ""))
            try:
                rec = call(g["path"], md5, g["hub"], cut)
                if cut and any(judge(ext, fn, rec["invoices"], pages, q)["status"] != "PASS" for _, ext, fn, q in need):
                    print("    the 2-3 page read did not pass -> sending the full PDF")
                    first = rec
                    rec = call(g["path"], md5, g["hub"])
                    rec["slice_try"] = {k: first[k] for k in ("pages_sent", "tokens", "cost_rs", "invoices")}
                    rec["cost_rs"] = round(rec["cost_rs"] + first["cost_rs"], 2)
                    save_raw(rec)
                    n["full_after_slice"] += 1
            except Exception as e:
                for row, *_ in need:
                    mark(row, None, "REVIEW", f"Gemini failed: {e}")
                    n["review"] += 1
                continue
            n["new_calls"] += 1
            n["sliced"] += bool(cut)
            n["cost_rs"] += rec["cost_rs"]
        text = rec.get("model") == "pdf-text"
        for row, ext, fn, q in need:
            j = judge(ext, fn, rec["invoices"], pages, q)
            m = match_invoice(ext, rec["invoices"])
            rescue = None if text else getattr(call, "retry", None)   # text answers passed 5 checks; an image won't do better
            if j["status"] == "REVIEW" and rescue and m is not None and not m.get("retried"):
                print(f"    {ext}: REVIEW -> retry with only the Tax Details page, upright")
                try:
                    a, cost = rescue(g["path"], m, q)
                except Exception as ex:          # not marked retried -> tried again next run
                    print(f"    [retry failed] {ext}: {ex}")
                else:
                    m["retried"] = True          # saved: a re-run never pays for this retry again
                    if a:
                        m["first_try"] = {k: m.get(k) for k in RESCUE_KEYS}
                        m.update({k: a[k] for k in RESCUE_KEYS}, retry_page=a["page"], retry_rot=a["rot"])
                    rec["cost_rs"] = round(rec.get("cost_rs", 0) + cost, 2)
                    n["cost_rs"] += cost
                    n["retried"] += 1
                    n["rescued"] += bool(a)
                    save_raw(rec)
                    j = judge(ext, fn, rec["invoices"], pages, q)
            if j["status"] == "PASS":
                ws.cell(row, COL_A5).value = j["a5"]
                ws.cell(row, COL_A18).value = j["a18"]
                ws.cell(row, COL_IGST).value = "y" if j["igst"] else None
                if j["date"] and not q:
                    ws.cell(row, COL_DATE).value = j["date"]      # printed date beats the CSV date (QR date set above)
                n["passed"] += 1
            else:
                n["review"] += 1
            if text:
                j["note"] = "; ".join(x for x in ("read from PDF text", j["note"]) if x)
            mark(row, j["printed"], j["status"], j["note"])
    return n


# ---------- Gemini_Results.xlsx: every answer vs what was finally entered ----------
def write_view():
    """Rebuild Gemini\\Gemini_Results.xlsx from Gemini\\raw\\*.json and every *Tax_Verification.csv
    (what generate_master_csvs put in the import files). Read-only for everything else."""
    import csv
    from openpyxl import Workbook
    from openpyxl.styles import Font
    recs = []
    for p in sorted(glob.glob(os.path.join(RAW, "*.json"))):
        with open(p, encoding="utf-8") as f:
            recs.append(json.load(f))
    if not recs:
        return None
    entered = {}   # invoice -> (taxable 5%, taxable 18%) from the newest import file that has it
    for p in sorted(glob.glob(os.path.join(BOOK_ROOT, "**", "*Tax_Verification.csv"), recursive=True),
                    key=os.path.getmtime):
        with open(p, newline="", encoding="utf-8") as f:
            for r in csv.DictReader(f):
                try:
                    entered[str(r["Invoice"]).strip()] = (float(r["Taxable5%"] or 0), float(r["Taxable18%"] or 0))
                except (KeyError, ValueError):
                    pass
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    ws.append(["Read on", "Hub", "PDF file", "Invoice (printed)", "Gemini 5%", "Gemini 18%", "IGST",
               "Printed tax", "Gemini check", "Entered 5%", "Entered 18%", "Gemini vs entered", "Read by"])
    s = dict(files=sum(r.get("model") != "pdf-text" for r in recs), text=sum(r.get("model") == "pdf-text" for r in recs), cost=sum(r.get("cost_rs", 0) for r in recs), answers=0, passed=0, same=0, diff=0, rescued=0)
    for rec in sorted(recs, key=lambda r: r["read_on"]):
        for e in rec["invoices"] or [{}]:
            inv = str(e.get("external_id", "")).strip()
            chk = reconcile(e.get("amount_5"), e.get("amount_18"), e.get("printed_total_tax"))[0] if e else "nothing read"
            if e.get("first_try"):
                chk += " (2nd try: tax page)"      # 1st whole-PDF answer failed; kept in the json as first_try
            t = entered.get(inv)
            verdict = ""
            if t:
                same = abs(t[0] - (e.get("amount_5") or 0)) <= 1 and abs(t[1] - (e.get("amount_18") or 0)) <= 1
                verdict = "SAME" if same else "DIFFERENT"
                if chk.startswith("PASS"):
                    s["same" if same else "diff"] += 1
            s["answers"] += bool(e)
            s["passed"] += chk.startswith("PASS")
            s["rescued"] += bool(e.get("first_try"))
            ws.append([rec["read_on"], rec.get("hub", ""), rec["file"], inv, e.get("amount_5"), e.get("amount_18"),
                       "y" if e.get("is_igst") else "", e.get("printed_total_tax"), chk,
                       t[0] if t else None, t[1] if t else None, verdict, rec.get("model", "")])
    for i, w in enumerate([16, 13, 52, 14, 13, 13, 5, 12, 12, 13, 13, 16, 16], start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    sm = wb.create_sheet("Summary", 0)
    for label, v in [("PDFs read by Gemini (each paid once)", s["files"]),
                     ("PDFs read from PDF text (free, 5 math checks)", s["text"]),
                     ("Total cost so far (Rs, approx)", round(s["cost"], 2)),
                     ("Invoices read", s["answers"]),
                     ("PASS (tax math matched the printed tax)", s["passed"]),
                     ("   ... of which on the 2nd try (tax page only)", s["rescued"]),
                     ("PASS answers later in an import file", s["same"] + s["diff"]),
                     ("   ... same as Gemini", s["same"]),
                     ("   ... DIFFERENT (you changed it - check why)", s["diff"])]:
        sm.append([label, v])
        sm.cell(sm.max_row, 1).font = Font(bold=True)
    sm.column_dimensions["A"].width = 46
    try:
        wb.save(VIEW)
    except PermissionError:
        print(f"  (close {os.path.basename(VIEW)} in Excel to refresh it - the answers are safe in Gemini\\raw)")
        return None
    return VIEW


def main():
    if "--selftest" in sys.argv:
        return _selftest()
    if "--view" in sys.argv:
        print("Refreshed:", write_view() or "nothing saved yet")
        return

    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    master = args[0] if args else MASTER
    root = os.path.dirname(os.path.abspath(master))   # PDFs live under <root>\<Hub>\PDFs
    if not os.path.exists(master):
        sys.exit(f'No entry workbook: {master}\nBuild it first: python build_master_workbook.py "<csv>"')
    try:
        open(master, "r+b").close()
    except PermissionError:
        sys.exit(f"Close {os.path.basename(master)} in Excel first - nothing was sent to Gemini.")

    from dotenv import load_dotenv
    load_dotenv(os.path.join(os.path.dirname(SCRIPT_DIR), ".env"))   # OMR\.env
    key = os.environ.get("GEMINI_API_KEY")
    call = make_caller(key) if key else None
    if not key:
        print("No GEMINI_API_KEY -> using SAVED Gemini answers only (no new calls). Fill the rest by hand.")

    from openpyxl import load_workbook
    wb = load_workbook(master)               # keep formulas (H..N), dropdown, hyperlinks intact
    n = fill_sheet(wb["Entry"], root, call)
    wb.save(master)
    view = write_view()

    print("-" * 64)
    print(f"Rows in Master_Entry ............... {n['rows']}")
    print(f"  QR: right PDF {n['qr_ok']} (date fixed {n['date_fixed']}), WRONG FILE {n['wrong_file']}"
          f" (Remark written -> Missing PDFs list), no QR {n['no_qr']}")
    print(f"  marked x (No Gemini) - skipped ... {n['marked']}")
    print(f"  already filled (you/Gemini) - skip {n['typed']}")
    print(f"  from saved answers (free) ........ {n['saved']}")
    print(f"  read from PDF text (free) ........ {n['text_read']}")
    print(f"  new Gemini calls ................. {n['new_calls']} PDF(s), about Rs {n['cost_rs']:.2f}")
    print(f"    sent as 2-3 pages first ........ {n['sliced']}  (full PDF needed after: {n['full_after_slice']})")
    print(f"  REVIEW retried on the tax page ... {n['retried']}  (now PASS: {n['rescued']})")
    print(f"  not read (no key / no PDF) ....... {n['not_read']}")
    print(f"Filled (checksum PASS) ............. {n['passed']}")
    print(f"Left for you (red REVIEW) .......... {n['review']}")
    print(f"Saved: {master}")
    if view:
        print(f"Gemini answers vs entered: {view}")
    print("Open it, fill any red REVIEW rows by hand (match 'Printed Tax (Gemini)'), then: python generate_master_csvs.py")


if __name__ == "__main__":
    main()
