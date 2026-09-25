"""Read Master_Entry.xlsx -> ONE combined Header CSV + Expenses CSV + Verification + one zip.

Usage:
    python generate_master_csvs.py
    python generate_master_csvs.py --no-vendor-select      # omit the Vendor Select address column
    python generate_master_csvs.py --skip=123,456          # External IDs already booked
    python generate_master_csvs.py --selfcheck             # run the tax-math self-test only
    python generate_master_csvs.py --accept=123,456        # book these although they fail the printed-tax check

MONEY CHECK: when Gemini read a row's printed total tax ("Printed Tax (Gemini)" column) and the
amounts in E/F give a different tax (more than Rs 1 off), NOTHING is written - fix those rows
first. (On 21-Sep four bills were booked with only the 5% amount this way.) If you checked the PDF
and Gemini's printed-tax read is the wrong one, pass --accept=<invoice>.
QR CHECK (same stop, same --accept): your taxable + tax must equal the e-invoice "QR Total" within
Rs 1, and a row whose QR Check says WRONG FILE must not have amounts.
Also refreshes the Missing PDFs list (rows left out of this batch) and Gemini\\Gemini_Results.xlsx.

Each Entry row carries its Hub (last column); the hub's vendor/GSTIN/address config is read
straight from build_entry_workbook.HUBS (no per-hub Config sheet needed). Tax type is derived
per row by process_entry_row exactly as the per-hub generator does.

Reads : <OUT_ROOT>\\Master_Entry.xlsx
Writes: <OUT_ROOT>\\1_ALL_Bills_Header.csv, 2_ALL_Bills_Expenses.csv,
        ALL_Tax_Verification.csv, ALL_Invoices.zip
"""
import csv, os, sys, zipfile
from openpyxl import load_workbook
from build_entry_workbook import HUBS, config_rows, OUT_ROOT, ENTRY_HEADERS, PRINTED_TAX, QR_CHECK, QR_TOTAL
from generate_import_csvs import HEADER_COLS, EXP_COLS, process_entry_row, fmt_date
import missing_tracker as mt

HUB_IDX = len(ENTRY_HEADERS)   # 0-based index of the Hub column (appended after N)


def tax_gap(a5, a18, printed):
    """Typed tax minus the invoice's printed tax, or None when there is nothing to compare."""
    if not isinstance(printed, (int, float)) or (a5 in (None, "") and a18 in (None, "")):
        return None
    try:
        return round(float(a5 or 0) * 0.05 + float(a18 or 0) * 0.18 - printed, 2)
    except (TypeError, ValueError):
        return None


def qr_total_gap(a5, a18, qr_total):
    """Your taxable + tax minus the e-invoice QR total, or None when there is nothing to compare."""
    if not isinstance(qr_total, (int, float)) or (a5 in (None, "") and a18 in (None, "")):
        return None
    try:
        a5, a18 = float(a5 or 0), float(a18 or 0)
    except (TypeError, ValueError):
        return None
    return round(a5 + a18 + round(a5 * 0.05 + a18 * 0.18, 2) - qr_total, 2)


def selfcheck():
    # intra hub (OMR: POS 33, GSTIN 33) and IGST hub (Pondicherry: POS 34, GSTIN 33)
    omr = dict(config_rows("omr"))
    h, e, v, w = process_entry_row("T1", "13/09/2026", "HUL", "x.pdf", 1000, 1000, "", omr, "33")
    assert not w and v[1] == "CGST+SGST" and v[5] == v[6] and str(v[5]) == "115.00", v
    assert str(v[9]) == "2.00" and str(v[10]) == "2228.00", v        # TDS, NetPayable
    assert len(e) == 2, e
    pny = dict(config_rows("pondicherry"))
    h, e, v, w = process_entry_row("T2", "13/09/2026", "HUL", "x.pdf", 1000, 1000, "", pny, "34")
    assert v[1] == "IGST" and str(v[7]) == "230.00" and str(v[10]) == "2228.00", v
    # money check: the real 21-Sep miss (only 5% typed) is caught; a right entry and no-Gemini rows pass
    assert tax_gap(102673.35, None, 175704.24) == -170570.57
    assert abs(tax_gap(151255.70, 0, 7562.84)) <= 1
    assert tax_gap(1000, 0, None) is None and tax_gap(None, None, 50.0) is None
    # QR check: real 9629078879 (Tumkur), QR total 5,35,184.89; a typo of Rs 4 in the 5% amount is caught
    assert abs(qr_total_gap(57043.02, 402787.93, 535184.89)) <= 1
    assert abs(qr_total_gap(57047.02, 402787.93, 535184.89)) > 1
    assert qr_total_gap(1000, 0, None) is None and qr_total_gap(None, "", 5.0) is None
    print("selfcheck OK")


def main():
    if "--selfcheck" in sys.argv:
        selfcheck(); return
    include_vs = "--no-vendor-select" not in sys.argv
    skip, accept = set(), set()
    for a in sys.argv:
        if a.startswith("--skip="):
            skip = {s.strip() for s in a.split("=", 1)[1].split(",") if s.strip()}
        if a.startswith("--accept="):
            accept = {s.strip() for s in a.split("=", 1)[1].split(",") if s.strip()}
    mt.assert_writable()                 # Missing PDFs list is updated at the end

    # per-Location config = the flat dict the old per-hub Config sheet used to hold
    cfg_by_loc, pos_by_loc = {}, {}
    for k in HUBS:
        C = dict(config_rows(k))
        loc = HUBS[k]["Location"]
        cfg_by_loc[loc] = C
        pos_by_loc[loc] = str(C["Place of Supply"]).split("-")[0].strip()

    entry_path = os.path.join(OUT_ROOT, "Master_Entry.xlsx")
    wb = load_workbook(entry_path, data_only=True)
    hdr = [c.value for c in wb["Entry"][1]]
    p_idx = hdr.index(PRINTED_TAX) if PRINTED_TAX in hdr else None
    q_idx = hdr.index(QR_CHECK) if QR_CHECK in hdr else None
    qt_idx = hdr.index(QR_TOTAL) if QR_TOTAL in hdr else None
    header_rows, exp_rows, verify_rows, warnings, blocked = [], [], [], [], []
    for r in wb["Entry"].iter_rows(min_row=2, values_only=True):
        inv = r[0]
        if inv is None:
            continue
        if str(inv) in skip:
            continue
        qr = str(r[q_idx] or "") if q_idx is not None and len(r) > q_idx else ""
        if (qr.startswith("WRONG FILE") and (r[4] not in (None, "") or r[5] not in (None, ""))
                and str(inv).strip() not in accept):
            blocked.append(f"  {inv}  ({r[HUB_IDX] if len(r) > HUB_IDX else ''}): amounts typed, but the QR says "
                           f"{qr} - this PDF is another invoice")
            continue
        qt = r[qt_idx] if qt_idx is not None and len(r) > qt_idx else None
        qgap = qr_total_gap(r[4], r[5], qt)
        if qgap is not None and abs(qgap) > 1 and str(inv).strip() not in accept:
            blocked.append(f"  {inv}  ({r[HUB_IDX] if len(r) > HUB_IDX else ''}): your taxable + tax is "
                           f"{qt + qgap:,.2f} but the QR total is {qt:,.2f}  (off by {qgap:,.2f})")
            continue
        gap = tax_gap(r[4], r[5], r[p_idx]) if p_idx is not None and len(r) > p_idx else None
        if gap is not None and abs(gap) > 1 and str(inv).strip() not in accept:
            blocked.append(f"  {inv}  ({r[HUB_IDX] if len(r) > HUB_IDX else ''}): your amounts give tax "
                           f"{r[p_idx] + gap:,.2f} but the PDF prints {r[p_idx]:,.2f}  (off by {gap:,.2f})")
            continue
        hub = str(r[HUB_IDX]).strip() if len(r) > HUB_IDX and r[HUB_IDX] else ""
        if hub not in cfg_by_loc:
            warnings.append(f"{inv}: unknown/blank Hub '{hub}' - row skipped")
            continue
        h, e, v, w = process_entry_row(inv, fmt_date(r[1]), r[2], r[3], r[4], r[5], r[6],
                                       cfg_by_loc[hub], pos_by_loc[hub])
        warnings += w
        if h is None:
            continue
        header_rows.append(h); exp_rows += e; verify_rows.append(v)

    if blocked:
        sys.exit("STOPPED - nothing written. These rows don't match their PDF:\n"
                 + "\n".join(blocked)
                 + "\nOpen each PDF, correct Amount_5% / Amount_18% (or clear them for a wrong file), save, and run again.\n"
                 "Only if you checked the PDF and it is right: "
                 "python generate_master_csvs.py --accept=<invoice>,<invoice>")
    assert header_rows, "No bills generated - did you fill amounts and save Master_Entry.xlsx?"

    header_cols = list(HEADER_COLS)
    if not include_vs:
        header_cols.remove("Vendor Select")
        for h in header_rows:
            h.pop("Vendor Select", None)

    h_path = os.path.join(OUT_ROOT, "1_ALL_Bills_Header.csv")
    e_path = os.path.join(OUT_ROOT, "2_ALL_Bills_Expenses.csv")
    v_path = os.path.join(OUT_ROOT, "ALL_Tax_Verification.csv")
    with open(h_path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=header_cols); w.writeheader(); w.writerows(header_rows)
    with open(e_path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=EXP_COLS); w.writeheader(); w.writerows(exp_rows)
    with open(v_path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["Invoice", "TaxType", "Taxable5%", "Taxable18%", "TotTaxable",
                    "CGST", "SGST", "IGST", "TotGST", "TDS194Q", "NetPayable"])
        for row in verify_rows:
            w.writerow([row[0], row[1]] + [str(v) for v in row[2:]])

    # Bundle every referenced PDF from all hubs' PDFs folders into one zip.
    wanted = [h["Attached file"] for h in header_rows if h.get("Attached file")]
    locs = {h["Location"] for h in header_rows}
    found = {}
    for loc in locs:
        for dp, _, files in os.walk(os.path.join(OUT_ROOT, loc, "PDFs")):
            for fn in files:
                if fn in wanted and fn not in found:
                    found[fn] = os.path.join(dp, fn)
    z_path = os.path.join(OUT_ROOT, "ALL_Invoices.zip")
    with zipfile.ZipFile(z_path, "w", zipfile.ZIP_DEFLATED) as z:
        for fn in wanted:
            if fn in found:
                z.write(found[fn], fn)
    for fn in wanted:
        if fn not in found:
            warnings.append(f"PDF not found under any hub's PDFs folder: {fn}")

    # Missing PDFs list: rows left out of this batch join it; rows in it are marked "in today's import"
    from mark_booked import chase_not_booked
    in_batch = {str(h["External ID"]).strip() for h in header_rows}
    chase = mt.load()
    left_out = chase_not_booked(entry_path, in_batch, chase)
    for inv in in_batch & set(chase):
        mt.upsert(chase, inv, status="in today's import - run mark_booked.py after NetSuite accepts")
    chase_path, open_n = mt.save(chase)
    from fill_tax_from_pdfs import write_view
    view = write_view()

    igst_n = sum(1 for v in verify_rows if v[1] == "IGST")
    per_hub = {}
    for h in header_rows:
        per_hub[h["Location"]] = per_hub.get(h["Location"], 0) + 1
    print(f"STEP 4 - Import files ready: {len(header_rows)} bills  ("
          + ", ".join(f"{k} {per_hub[k]}" for k in sorted(per_hub)) + ")")
    print(f"  Header  : {h_path}")
    print(f"  Expenses: {e_path}  ({len(exp_rows)} lines)")
    print(f"  Verify  : {v_path}   Tax split: {len(verify_rows) - igst_n} CGST+SGST, {igst_n} IGST")
    print(f"  Zip     : {z_path}  ({sum(1 for fn in wanted if fn in found)}/{len(wanted)} PDFs)")
    print(f"  Vendor Select (address): {'INCLUDED' if include_vs else 'OMITTED'}")
    print(f"  Printed-tax + QR check: every row with a Gemini printed tax / QR total matches (accepted by you: {len(accept)})")
    print(f"  Left out of this batch: {left_out} -> Missing PDFs list ({open_n} open): {chase_path}")
    if view:
        print(f"  Gemini answers vs entered: {view}")
    for w_ in warnings:
        print("WARN:", w_)
    print("Next: NetSuite import (zip FIRST), then: python mark_booked.py")


if __name__ == "__main__":
    main()
