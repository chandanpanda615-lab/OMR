"""Record invoices as BOOKED after a successful NetSuite import, so they can never
be re-downloaded, re-listed, or re-booked by accident.

    python mark_booked.py                       # MASTER flow: reads 1_ALL_Bills_Header.csv
    python mark_booked.py "Mysore Road" Byrathi # per-hub flow: reads 1_<HUB>_Bills_Header.csv

Both modes append External IDs to 5_NetSuite_Booking\\_booked.csv
(invoice_no, hub, invoice_date, marked_on), deduped. build_master_workbook.py,
download_hul_pdfs.py and download_from_links.py all read this ledger and skip
anything already in it.

MASTER mode also MOVES each booked PDF from <hub>\\PDFs\\... to <hub>\\Booked\\...,
so the PDFs folder always shows only invoices still waiting to be booked, and updates the
chase list (MISSING_PDFS.xlsx): booked invoices leave it, listed-but-not-booked rows join it
(Remark typed -> "wrong file", that PDF remembered so it is never listed again).

Self-check: run it twice -> the second run reports every invoice as "already in
ledger" and adds 0 rows. That is the dedupe test, on real data.
"""
import csv, os, shutil, sys
from datetime import date

OUT_ROOT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "5_NetSuite_Booking")
LEDGER = os.path.join(OUT_ROOT, "_booked.csv")


def load_booked():
    booked = set()
    if os.path.exists(LEDGER):
        with open(LEDGER, newline="", encoding="utf-8") as f:
            for row in csv.DictReader(f):
                booked.add(str(row["invoice_no"]).strip())
    return booked


def chase_not_booked(entry_path, done, chase):
    """Every Master_Entry row NOT in `done` goes on the chase list: with a Remark -> "wrong file"
    (its PDF remembered as bad_file, so the same file is never listed again); without -> "left blank"
    (+ Gemini's note, e.g. "not printed in PDF - wrong file?"). Returns how many rows were chased."""
    from openpyxl import load_workbook
    import missing_tracker as mt
    it = load_workbook(entry_path, read_only=True, data_only=True)["Entry"].iter_rows(values_only=True)
    col = {str(h): i for i, h in enumerate(next(it)) if h}

    def get(r, name):
        i = col.get(name)
        return "" if i is None or i >= len(r) or r[i] is None else str(r[i]).strip()

    n = 0
    for r in it:
        inv = get(r, "External ID")
        if not inv or inv in done:
            continue
        common = dict(hub=get(r, "Hub"), brand=get(r, "Brand"), invoice_date=r[1])
        if get(r, "Remark") and get(r, "Remark") != mt.VERIFY:   # untouched VERIFY = not checked yet
            mt.upsert(chase, inv, status="wrong file - ask hub to re-upload in CDMS",
                      remark=get(r, "Remark"), bad_file=get(r, "Attached file"), **common)
        else:
            note = get(r, "Gemini note")
            mt.upsert(chase, inv, status="left blank in Master_Entry - not booked"
                      + (" (new PDF not checked yet)" if get(r, "Remark") else "")
                      + (f" (Gemini: {note})" if note else ""), **common)
        n += 1
    return n


def move_to_booked(hub, fn):
    """Move a booked PDF from <hub>\\PDFs\\... to <hub>\\Booked\\. Returns
    True if it now lives under Booked (moved or already there), False if not found."""
    if not (hub and fn):
        return False
    dest_dir = os.path.join(OUT_ROOT, hub, "Booked")
    dest = os.path.join(dest_dir, fn)
    if os.path.exists(dest):
        return True                      # already marked in a previous run
    for dp, _, files in os.walk(os.path.join(OUT_ROOT, hub, "PDFs")):
        if fn in files:
            os.makedirs(dest_dir, exist_ok=True)
            shutil.move(os.path.join(dp, fn), dest)
            return True
    return False


def main():
    hubs = [a for a in sys.argv[1:] if not a.startswith("-")]
    master = not hubs                    # no hub names -> master (ALL) flow
    if master:
        import missing_tracker
        missing_tracker.assert_writable()   # stop before the ledger is touched, not half-way
    booked = load_booked()
    new_rows, already, moved, no_pdf = [], 0, 0, 0
    today = date.today().isoformat()

    if master:
        p = os.path.join(OUT_ROOT, "1_ALL_Bills_Header.csv")
        assert os.path.exists(p), f"missing: {p} — run generate_master_csvs.py first"
        with open(p, newline="", encoding="utf-8") as f:
            for row in csv.DictReader(f):
                inv = str(row["External ID"]).strip()
                if inv in booked:
                    already += 1
                    continue
                booked.add(inv)
                hub = str(row.get("Location", "")).strip()
                new_rows.append([inv, hub, row.get("Invoice Date", ""), today])
                if move_to_booked(hub, str(row.get("Attached file", "")).strip()):
                    moved += 1
                else:
                    no_pdf += 1
    else:
        for h in hubs:
            p = os.path.join(OUT_ROOT, h, f"1_{h}_Bills_Header.csv")
            assert os.path.exists(p), f"missing: {p} — generate its import CSV first"
            with open(p, newline="", encoding="utf-8") as f:
                for row in csv.DictReader(f):
                    inv = str(row["External ID"]).strip()
                    if inv in booked:
                        already += 1
                        continue
                    booked.add(inv)
                    new_rows.append([inv, h, row.get("Invoice Date", ""), today])

    first_write = not os.path.exists(LEDGER)
    with open(LEDGER, "a", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        if first_write:
            w.writerow(["invoice_no", "hub", "invoice_date", "marked_on"])
        w.writerows(new_rows)
    print(f"Marked booked: {len(new_rows)} new  (already in ledger: {already})")
    if master:
        print(f"PDFs moved to Booked: {moved}   (PDF not found on disk: {no_pdf})")
    print(f"Ledger: {LEDGER}  (total {len(booked)} invoices)")

    if master:   # booked rows leave the chase list; listed-but-not-booked rows join it
        import missing_tracker as mt
        chase = mt.load()
        entry_path = os.path.join(OUT_ROOT, "Master_Entry.xlsx")
        n = chase_not_booked(entry_path, booked, chase) if os.path.exists(entry_path) else 0
        path, left = mt.save(chase)
        print(f"Listed but not booked: {n} -> Missing PDFs list  |  now {left} open: {path}")


if __name__ == "__main__":
    main()
